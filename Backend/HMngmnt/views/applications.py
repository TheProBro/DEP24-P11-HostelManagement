from django.http import JsonResponse
from django.http import HttpResponse
from django.views.decorators.csrf import csrf_exempt

from ..models import CustomUser, Faculty, Application, Application_Final, Hostel, SavedMappings
from ..decorators import staff_required, token_required
from ..email import send, templates
from ..helpers import handle_file_attachment

import json
import os

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

@csrf_exempt
@staff_required
def view_applications(request):
    user=request.new_param
    # print('user:', user)
    faculty=Faculty.objects.get(faculty=CustomUser.objects.get(pk=user.get('id')))
    if not faculty.is_hod:
        print('faculty:', faculty)
        applications = Application.objects.filter(faculty=faculty).select_related('student', 'faculty')
        application_list=[{
            'application_id': application.application_id,
            'student': application.student.name,
            'affiliation': application.affiliation,
            'faculty': application.faculty.faculty.name,
            'status': application.status
        } for application in applications]
        print('application_list:', application_list)
        return JsonResponse({'message': 'Staff page', 'data': {'own': application_list}})
    else:
        print("HOD here")
        # applications for hod approval
        applications1=Application.objects.filter(status='Pending HOD Approval', faculty__department=faculty.department).select_related('student', 'faculty')
        #applications for own approval
        applications2 = Application.objects.filter(faculty=faculty).select_related('student', 'faculty')

        application_list1=[{
            'application_id': application.application_id,
            'student': application.student.name,
            'affiliation': application.affiliation,
            'faculty': application.faculty.faculty.name,
            'status': application.status
        } for application in applications1]
        application_list2=[{
            'application_id': application.application_id,
            'student': application.student.name,
            'affiliation': application.affiliation,
            'faculty': application.faculty.faculty.name,
            'status': application.status
        } for application in applications2]
        # print('application_list:', application_list)
        # return JsonResponse({'message': 'Staff page', 'data': application_list})
        return JsonResponse({'message': 'Staff page', 'data': {'hod': application_list1, 'own': application_list2}})

@csrf_exempt
@staff_required
def view_final_applications(request):
    user=request.new_param
    user=CustomUser.objects.get(pk=user.get('id'))
    hostel=user.caretaker.hostel
    print('hostel:', hostel)
    applications=Application_Final.objects.select_related('application', 'application__student', 'application__faculty').filter(hostel=hostel)
    application_list = [
        {
            'application_id': application.application.application_id,
            'student': application.application.student.name,
            'affiliation': application.application.affiliation,
            'faculty': application.application.faculty.faculty.name,
            'status': application.application.status
        }
        for application in applications
    ]
    print('application_list:', application_list)
    # return JsonResponse({'message': 'Staff page', 'data': {'own': []}})
    return JsonResponse({'message': 'Staff page', 'data': {'own':application_list}})    


@csrf_exempt
@staff_required
def update_application(request):
    try:
        # print(request.body.decode('utf-8'))
        data=json.loads(request.body.decode('utf-8'))
        data=data['updatedSelectedOptions']
        # print('here',data)
        # print(request.path)
        for application_id, status in data.items():
            # print('application_id:', application_id)
            # print('status:', status['hostel'])
            application=Application.objects.get(application_id=int(application_id))
            application.status=status['value']
            print('application:', application)
            if 'Rejected' in status['value']:
                application.comments=status['comments']
            elif status['value']=='Pending Caretaker Action':
                # Application_Final(application=application, hostel=Hostel.objects.get(hostel_name=status['hostel'])).save()
                hostel=str(status['hostel']).split(' ')[0]
                hostel=Hostel.objects.get(hostel_name=hostel)
                # print('here')
                # print('hostel:', hostel)
                new_app=Application_Final(application=application, hostel=hostel)
                new_app.save()
            application.save()
        return JsonResponse({'message': 'Applications updated successfully.'})
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=400)

@csrf_exempt
@staff_required
def send_email(request):
    recipient=request.GET.get('recipient')
    template=int(request.GET.get('template'))
    if(template==0):
        id=int(request.GET.get('id')) or None
        application=Application.objects.get(application_id=id)
        application.status='Pending Payment'
        application.save()
    template=templates[int(template)]
    send(template, [recipient])
    return JsonResponse({'message': 'Email sent successfully.'})

@csrf_exempt
# @staff_required
def generate_pdf(request):
    application_id = request.GET.get('application_id') or None
    allotments=request.GET.get('allotments') or None
    if application_id:
        application = Application.objects.get(application_id=application_id)

        # Create a list to hold the data for the PDF
        data = []

        # Add the text fields to the data list
        data.append(['Student Name:', application.student.name])
        data.append(['Affiliation:', application.affiliation])
        data.append(['Faculty:', application.faculty.faculty.name])
        # data.append(['Status:', application.status])
        data.append(['Address:', application.address])
        data.append(['Arrival Date:', application.arrival])
        data.append(['Departure Date:', application.departure])

        # Create a PDF document
        doc = SimpleDocTemplate("./documents/datagen.pdf", pagesize=letter)

        # Create a list to hold the flowables (elements) of the document
        elements = []

        # Create a style for the text fields
        styles = getSampleStyleSheet()
        style = styles["Normal"]

        # Add the text fields to the document
        for field in data:
            text = f"{field[0]} {field[1]}"
            p = Paragraph(text, style)
            elements.append(p)
            elements.append(Spacer(1, 12))

        # Build the document and save it
        doc.build(elements)

        # Send the PDF file as a response to the frontend
        with open("./documents/datagen.pdf", "rb") as f:
            response = HttpResponse(f.read(), content_type="application/pdf")
            response["Content-Disposition"] = "inline; filename=output.pdf"
        
        # delete the file
        os.remove("./documents/datagen.pdf")
        return response
    elif allotments:
        print(allotments)
        # generate a sheet of current mappings
        mp1=SavedMappings.objects.get(name=allotments)
        # mp2=SavedMappings.objects.get(name='Current Girls')
        workbook = Workbook()
        worksheet = workbook.active
        row = 1
        col = 1
        for data_row in mp1.mapping:
            for cell in data_row:
                worksheet.cell(row=row, column=col, value=cell)
                col += 1
            col = 1
            row += 1
        worksheet.cell(row=row, column=1, value="Wing Room Capacities")
        row += 1
        # Write each key-value pair to the Excel sheet
        for wing, capacity in mp1.wing_room_capacities.items():
            worksheet.cell(row=row, column=1, value=wing)
            worksheet.cell(row=row, column=2, value=capacity)
            row += 1
        worksheet.cell(row=row, column=1, value="Batch Strengths")
        row += 1
        # Write each key-value pair to the Excel sheet
        for batch, strength in mp1.batch_strengths.items():
            worksheet.cell(row=row, column=1, value=batch)
            worksheet.cell(row=row, column=2, value=strength)
            row += 1
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = f'attachment; filename="Boys_data.xlsx"'
        workbook.save(response)

    return response

@csrf_exempt
@token_required
def unallocate_room(request):
    id=request.GET.get('application_id')
    app=Application.objects.get(application_id=id)
    app.status='Done'
    app_final=Application_Final.objects.get(application=app)
    app_final.delete()
    app.save()
    return JsonResponse({'message': 'Room unallocated successfully'})

@csrf_exempt
@token_required
def get_application_status(request):
    user=request.new_param
    application=Application.objects.filter(student__id=user.get('id')).select_related('student', 'faculty').first()
    if application is None:
        return JsonResponse({'message': 'Student page', 'data': None})
    applicationX={
        'application_id': application.application_id,
        'student': application.student.name,
        'affiliation': application.affiliation,
        'faculty': application.faculty.faculty.name,
        'status': application.status,
        'address': application.address,
        'arrival': application.arrival,
        'departure': application.departure,
        'phone': application.phone,
        'email': application.student.email,
        'instiId': handle_file_attachment(application.instiId.path),
        'letter': handle_file_attachment(application.letter.path),
        'payment_proof':handle_file_attachment(application.payment_proof),
        'transaction_id':application.payment_id,
    }
    if "Rejected" in application.status:
        applicationX['comments']=application.comments
    # print('here: ', application.letter.path, application.instiId.path)
    return JsonResponse({'message': 'Student page', 'data': applicationX})

@csrf_exempt
@token_required
def get_application(request, id):
    application=Application.objects.select_related('faculty', 'student').get(application_id=id)
    applicationX={
        'student_name': application.student.name,
        'application_id': application.application_id,
        'affiliation': application.affiliation,
        'faculty': application.faculty.faculty.name,
        'status': application.status,
        'student_contact_number': application.phone,
        'address': application.address,
        'arrival_date': application.arrival,
        'departure_date': application.departure,
        'institute_id': handle_file_attachment(application.instiId.path),
        'institute_letter': handle_file_attachment(application.letter.path),
        'student_email': application.student.email,
        'payment_proof': handle_file_attachment(application.payment_proof.path) if application.payment_proof else None,
        'payment_id': application.payment_id if application.payment_id else None
    }
    # print('here:', application.instiId.path, application.letter.path)
    return JsonResponse({'message': 'Student page', 'data': applicationX})

@csrf_exempt
@token_required
def internship(request):
    facultyE= request.POST.get('facultyEmail')
    student= request.POST.get('studentEmail')
    gender= request.POST.get('gender')
    correction= request.POST.get('correction')
    if correction is not None:
        application=Application.objects.get(application_id=correction)
        print(request.POST)
        for key, val in request.POST.items():
            print(key, val)
            if key=='correction' or key=='studentName' or key=='studentEmail' or val=='undefined':
                continue
            if val:
                if key=='gender':
                    application.student.gender=val
                    application.student.save()
                elif key=='affiliation':
                    application.affiliation=val
                elif key=='address':
                    application.address=val
                elif key=='contactNumber':
                    application.phone=val
                elif key=='facultyEmail':
                    application.faculty=Faculty.objects.get(faculty__email=val)
                elif key=='arrivalDate':
                    application.arrival=val
                elif key=='departureDate':
                    application.departure=val

        if request.FILES.get('instituteID'):
            application.instiId=request.FILES.get('instituteID')
        if request.FILES.get('instituteLetter'):
            application.letter=request.FILES.get('instituteLetter')
        application.comments='None'
        if "Faculty" in application.status:
            application.status='Pending Faculty Approval'
        elif "HOD" in application.status:
            application.status='Pending HOD Approval'
        elif "Admin" in application.status:
            application.status='Pending Admin Approval'
        application.save()

        # print(application, 'hahaha', request.POST)
        return HttpResponse("Application corrected successfully", status=200)
    if CustomUser.objects.get(email=student) is None:
        return HttpResponse("Student not found", status=300)
    if CustomUser.objects.get(email=facultyE) is None:
        return HttpResponse("Faculty not found", status=300)
    application = Application(
        student= CustomUser.objects.get(email=student),
        affiliation= request.POST.get('affiliation'),
        address= request.POST.get('address'),
        phone= request.POST.get('contactNumber'),
        faculty= Faculty.objects.get(faculty__email=facultyE),
        arrival= request.POST.get('arrivalDate'),
        departure= request.POST.get('departureDate'),
        instiId= request.FILES.get('instituteID'),
        letter= request.FILES.get('instituteLetter'),
    )
    print("hello")
    #update gender in user
    user=CustomUser.objects.get(email=student)
    user.gender=gender
    user.save()

    application.save()
    return HttpResponse("Application submitted successfully", status=200)


@csrf_exempt
@token_required
def update_payment(request):
    if request.method=='POST':
        # data=json.loads(request.body)
        application=Application.objects.get(application_id=request.POST.get('application_id'))
        application.payment_proof=request.FILES.get('payment_proof')
        application.payment_id=request.POST.get('payment_id')
        application.status='Payment Proof Uploaded'
        application.save()
        return JsonResponse({'message': 'Payment verified successfully'})
    return JsonResponse({'error': 'Invalid request'}, status=400)
